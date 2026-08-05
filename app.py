import http.server
import socketserver
import subprocess
import json
import os
import threading
import logging
import urllib.parse
import uuid
from datetime import datetime

import schedule
import time

from scraper import (
    canonical_product_key,
    enrich_record,
    load_records,
    merge_record,
    normalize_record,
    seller_to_legacy_text,
    write_records,
)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

PORT = int(os.getenv("TRAP_MONITOR_PORT", "8000"))
DIRECTORY = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(DIRECTORY, "data", "scraped_data.json")
LAST_RUN_FILE = os.path.join(DIRECTORY, "data", "last_run.json")
os.makedirs(os.path.join(DIRECTORY, "logs"), exist_ok=True)
DB_LOCK = threading.Lock()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(os.path.join(DIRECTORY, "logs", "server.log"), encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger(__name__)


def _load_all_records() -> list:
    with DB_LOCK:
        records = load_records()
        # 기존 데이터를 읽는 순간 새 데이터 모델과 중복 정규화를 반영합니다.
        return write_records(records)


def _save_all_records(records: list) -> list:
    with DB_LOCK:
        return write_records(records)


def _seller_from_payload(payload: dict, fallback_url: str = "") -> dict:
    seller = payload.get("seller_info") if isinstance(payload.get("seller_info"), dict) else {}
    return {
        "name": str(seller.get("name") or payload.get("seller_name") or "").strip(),
        "store_name": str(seller.get("store_name") or payload.get("seller_store_name") or "").strip(),
        "representative": str(seller.get("representative") or payload.get("seller_rep") or "").strip(),
        "address": str(seller.get("address") or payload.get("seller_address") or "").strip(),
        "phone": str(seller.get("phone") or payload.get("seller_phone") or "").strip(),
        "business_number": str(seller.get("business_number") or payload.get("business_number") or "").strip(),
        "source_url": str(seller.get("source_url") or fallback_url or "").strip(),
        "confidence": str(seller.get("confidence") or "manual").strip(),
    }


def _run_scraper_subprocess():
    """스크래퍼를 서브프로세스로 실행합니다."""
    venv_python = os.path.join(DIRECTORY, ".venv", "bin", "python")
    if not os.path.exists(venv_python):
        venv_python = os.path.join(DIRECTORY, ".venv", "Scripts", "python.exe")
    if not os.path.exists(venv_python):
        venv_python = "python3"

    scraper_script = os.path.join(DIRECTORY, "scraper.py")
    return subprocess.run(
        [venv_python, scraper_script],
        capture_output=True,
        text=True,
        cwd=DIRECTORY,
    )


def scheduled_scrape():
    logger.info("스케줄 자동 실행: 스크래핑 시작")
    result = _run_scraper_subprocess()
    if result.returncode == 0:
        logger.info("스케줄 자동 실행 완료")
    else:
        logger.error(f"스케줄 자동 실행 실패:\n{result.stderr}")


def _next_run_4th() -> str:
    """다음 매월 4일 날짜 문자열 반환"""
    now = datetime.now()
    if now.day < 4:
        next_run = now.replace(day=4)
    else:
        month = now.month + 1 if now.month < 12 else 1
        year = now.year if now.month < 12 else now.year + 1
        next_run = now.replace(year=year, month=month, day=4)
    return next_run.strftime("%Y-%m-%d")


def _start_scheduler():
    """매월 4일 00:00에 자동 스캔 스케줄 등록 후 백그라운드 실행"""
    schedule.every().day.at("00:00").do(lambda: (
        scheduled_scrape() if datetime.now().day == 4 else None
    ))
    logger.info(f"스케줄러 시작. 다음 자동 실행 예정일: 매월 4일 (다음: {_next_run_4th()})")

    def _loop():
        while True:
            schedule.run_pending()
            time.sleep(60)

    t = threading.Thread(target=_loop, daemon=True)
    t.start()


def _generate_excel(month_str: str, report_data: list) -> bytes:
    """
    month_str: "2026년 6월" 형태
    report_data: [{platform, item, url, seller(파싱됨), date}, ...]
    양식: 포획도구 판매처 목록 및 보고 양식
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "포획도구 판매처 목록 및 보고 양식"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    title_font = Font(bold=True, size=13)
    header_font = Font(bold=True, size=10)

    # 열 너비 설정
    col_widths = {
        "A": 6,   # 연번
        "B": 14,  # 플랫폼명
        "C": 14,  # 품목
        "D": 45,  # 게시물URL
        "E": 16,  # 업체명
        "F": 10,  # 대표자
        "G": 30,  # 주소
        "H": 14,  # 연락처
        "I": 16,  # 유역(지방)청
        "J": 12,  # 점검일
        "K": 14,  # 점검기관
        "L": 20,  # 조치사항
        "M": 14,  # 게시물삭제여부
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 행 1: 제목
    ws.merge_cells("A1:M1")
    ws["A1"] = f"{month_str} 포획도구 판매처 목록 및 보고 양식"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws["A1"].fill = PatternFill("solid", fgColor="BDD7EE")
    ws.row_dimensions[1].height = 28

    # 행 2: 대분류 헤더
    headers_row2 = {
        "A": "연번", "B": "플랫폼명", "C": "품목", "D": "게시물URL",
        "I": "점검결과",
    }
    ws.merge_cells("A2:A3")
    ws.merge_cells("B2:B3")
    ws.merge_cells("C2:C3")
    ws.merge_cells("D2:D3")
    ws.merge_cells("E2:H2")  # 판매자 인적사항
    ws.merge_cells("I2:M2")  # 점검결과

    for col, val in [("A", "연번"), ("B", "플랫폼명"), ("C", "품목"), ("D", "게시물URL"),
                     ("E", "판매자 인적사항"), ("I", "점검결과")]:
        cell = ws[f"{col}2"]
        cell.value = val
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    # 행 3: 소분류 헤더
    sub_headers = {
        "E": "업체명", "F": "대표자", "G": "주소", "H": "연락처",
        "I": "유역(지방)청\n(시·군·구)", "J": "점검일", "K": "점검기관",
        "L": "조치사항\n(조치일)", "M": "게시물\n삭제 여부",
    }
    for col, val in sub_headers.items():
        cell = ws[f"{col}3"]
        cell.value = val
        cell.font = header_font
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    # 병합된 헤더 셀에도 border 적용
    for col in ["A", "B", "C", "D"]:
        for row in [2, 3]:
            ws[f"{col}{row}"].border = border
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 32

    # 데이터 행
    for i, record in enumerate(report_data, start=1):
        row = i + 3

        seller = record.get("seller", "")
        s_name, s_rep, s_addr, s_phone = "-", "-", "-", "-"
        if "[업체명]" in seller:
            for part in seller.split("\n"):
                part = part.strip()
                if part.startswith("[업체명]"):
                    s_name = part.replace("[업체명]", "").strip()
                elif part.startswith("[대표자]"):
                    s_rep = part.replace("[대표자]", "").strip()
                elif part.startswith("[주소]"):
                    s_addr = part.replace("[주소]", "").strip()
                elif part.startswith("[연락처]"):
                    s_phone = part.replace("[연락처]", "").strip()
        else:
            s_name = seller

        values = [
            i,                          # A 연번
            record.get("platform", ""), # B 플랫폼명
            record.get("item", ""),     # C 품목
            record.get("url", ""),      # D URL
            s_name,                     # E 업체명
            s_rep,                      # F 대표자
            s_addr,                     # G 주소
            s_phone,                    # H 연락처
            "",                         # I 유역청
            "",                         # J 점검일
            "",                         # K 점검기관
            "",                         # L 조치사항
            "",                         # M 삭제여부
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = border
            cell.alignment = center if col_idx != 4 else left_wrap
            if col_idx == 1:
                cell.font = Font(bold=True, size=9)
            elif col_idx == 4:
                cell.font = Font(color="0563C1", underline="single", size=8)
            else:
                cell.font = Font(size=9)

        ws.row_dimensions[row].height = 18

    # 바이트로 변환
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Handler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=DIRECTORY, **kwargs)

    def log_message(self, format, *args):
        pass  # 기본 HTTP 로그 억제

    def _send_json(self, code: int, body):
        data = json.dumps(body, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _read_json_body(self) -> dict:
        content_length = int(self.headers.get("Content-Length", 0))
        if not content_length:
            return {}
        return json.loads(self.rfile.read(content_length).decode("utf-8"))

    def do_GET(self):
        route = urllib.parse.urlsplit(self.path).path
        if route == "/api/records":
            try:
                records = [record for record in _load_all_records() if record.get("status") != "deleted"]
                self._send_json(200, {"status": "success", "records": records})
            except Exception as exc:
                logger.exception("데이터 조회 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})
            return
        super().do_GET()

    def do_POST(self):
        route = urllib.parse.urlsplit(self.path).path

        if route == "/run-scraper":
            try:
                result = _run_scraper_subprocess()
                self._send_json(200, {
                    "status": "success" if result.returncode == 0 else "error",
                    "output": result.stdout,
                    "error": result.stderr,
                })
            except Exception as e:
                self._send_json(500, {"status": "error", "message": str(e)})

        elif route == "/api/records/import":
            try:
                body = self._read_json_body()
                incoming_records = body.get("records", [])
                if not isinstance(incoming_records, list):
                    self._send_json(400, {"status": "error", "message": "records 배열이 필요합니다."})
                    return

                records = _load_all_records()
                by_key = {(record.get("canonical_key") or record["id"]): record for record in records}
                imported = 0
                for raw in incoming_records:
                    # 초기 버전의 가상 예시 데이터는 실제 적발 자료로 가져오지 않습니다.
                    if raw.get("url") == "https://daangn.com/articles/12345":
                        continue
                    raw = dict(raw)
                    if not raw.get("source"):
                        legacy_id = str(raw.get("id", ""))
                        raw["source"] = "search" if legacy_id.startswith(("google_", "naver_")) else "manual"
                    raw.setdefault("status", "confirmed" if raw.get("source") == "manual" else "candidate")
                    record = normalize_record(raw)
                    key = record.get("canonical_key") or record["id"]
                    if key in by_key:
                        by_key[key] = merge_record(by_key[key], record)
                    else:
                        by_key[key] = record
                        imported += 1
                _save_all_records(list(by_key.values()))
                self._send_json(200, {"status": "success", "imported": imported})
            except Exception as exc:
                logger.exception("기존 브라우저 데이터 가져오기 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})

        elif route == "/api/records/manual":
            try:
                body = self._read_json_body()
                url = str(body.get("url", "")).strip()
                if not url.startswith(("http://", "https://")):
                    self._send_json(400, {"status": "error", "message": "올바른 게시물 URL이 필요합니다."})
                    return
                seller_info = _seller_from_payload(body, url)
                now = datetime.now().isoformat(timespec="seconds")
                record = normalize_record({
                    "id": "manual_" + uuid.uuid4().hex[:16],
                    "source": "manual",
                    "status": "confirmed",
                    "platform": body.get("platform", ""),
                    "item": body.get("item", "기타 불법 엽구"),
                    "query": body.get("item", ""),
                    "title": body.get("title", "수동 등록"),
                    "url": url,
                    "seller_info": seller_info,
                    "seller": seller_to_legacy_text(seller_info),
                    "date": now,
                    "first_seen": now,
                    "last_seen": now,
                    "reviewed_at": now,
                    "review_note": body.get("review_note", "수동 등록"),
                })
                records = _load_all_records()
                key = record.get("canonical_key") or canonical_product_key(url) or record["id"]
                found = False
                for idx, existing in enumerate(records):
                    if (existing.get("canonical_key") or existing["id"]) == key:
                        records[idx] = merge_record(existing, record)
                        records[idx]["status"] = "confirmed"
                        found = True
                        break
                if not found:
                    records.append(record)
                _save_all_records(records)
                self._send_json(200, {"status": "success", "record": record})
            except Exception as exc:
                logger.exception("수동 등록 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})

        elif route == "/api/records/review":
            try:
                body = self._read_json_body()
                record_id = str(body.get("id", ""))
                status = str(body.get("status", ""))
                if status not in {"candidate", "confirmed", "rejected", "auto_rejected"}:
                    self._send_json(400, {"status": "error", "message": "올바른 검토 상태가 필요합니다."})
                    return
                records = _load_all_records()
                target = next((record for record in records if record.get("id") == record_id), None)
                if not target:
                    self._send_json(404, {"status": "error", "message": "대상을 찾을 수 없습니다."})
                    return
                target["status"] = status
                target["review_note"] = str(body.get("review_note", target.get("review_note", ""))).strip()
                target["reviewed_at"] = datetime.now().isoformat(timespec="seconds")
                if body.get("title"):
                    target["title"] = str(body["title"]).strip()
                if body.get("item"):
                    target["item"] = str(body["item"]).strip()
                if isinstance(body.get("seller_info"), dict):
                    target["seller_info"] = _seller_from_payload(body, target.get("url", ""))
                    target["seller"] = seller_to_legacy_text(target["seller_info"])
                _save_all_records(records)
                self._send_json(200, {"status": "success", "record": normalize_record(target)})
            except Exception as exc:
                logger.exception("검토 저장 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})

        elif route == "/api/records/enrich":
            try:
                body = self._read_json_body()
                record_id = str(body.get("id", ""))
                records = _load_all_records()
                target_index = next((idx for idx, record in enumerate(records) if record.get("id") == record_id), None)
                if target_index is None:
                    self._send_json(404, {"status": "error", "message": "대상을 찾을 수 없습니다."})
                    return
                enriched = enrich_record(records[target_index])
                records[target_index] = merge_record(records[target_index], enriched)
                _save_all_records(records)
                self._send_json(200, {"status": "success", "record": records[target_index]})
            except Exception as exc:
                logger.exception("상품·판매자정보 자동 확인 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})

        elif route == "/api/records/delete":
            try:
                body = self._read_json_body()
                record_id = str(body.get("id", ""))
                records = _load_all_records()
                target = next((record for record in records if record.get("id") == record_id), None)
                if not target:
                    self._send_json(404, {"status": "error", "message": "대상을 찾을 수 없습니다."})
                    return
                target["status"] = "deleted"
                target["reviewed_at"] = datetime.now().isoformat(timespec="seconds")
                target["review_note"] = str(body.get("review_note", "사용자 삭제"))
                _save_all_records(records)
                self._send_json(200, {"status": "success"})
            except Exception as exc:
                logger.exception("삭제 처리 오류")
                self._send_json(500, {"status": "error", "message": str(exc)})

        elif route == "/reset-data":
            try:
                if os.path.exists(DB_FILE):
                    with open(DB_FILE, "w", encoding="utf-8") as f:
                        json.dump([], f, ensure_ascii=False)
                self._send_json(200, {"status": "success"})
            except Exception as e:
                self._send_json(500, {"status": "error", "message": str(e)})

        elif route == "/generate-report":
            if not OPENPYXL_AVAILABLE:
                self._send_json(500, {"status": "error", "message": "openpyxl 패키지가 설치되지 않았습니다."})
                return

            body = self._read_json_body()

            month_str = body.get("month", "")
            report_data = body.get("data", [])

            if not month_str or not report_data:
                self._send_json(400, {"status": "error", "message": "month와 data가 필요합니다."})
                return

            try:
                xlsx_bytes = _generate_excel(month_str, report_data)
                filename = f"포획도구_판매처목록_{month_str.replace(' ', '_').replace('년', 'y').replace('월', 'm')}.xlsx"
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename*=UTF-8\'\'{urllib.parse.quote(filename)}')
                self.send_header("Content-Length", str(len(xlsx_bytes)))
                self.end_headers()
                self.wfile.write(xlsx_bytes)
            except Exception as e:
                logger.error(f"엑셀 생성 오류: {e}")
                self._send_json(500, {"status": "error", "message": str(e)})

        elif route == "/scheduler-status":
            try:
                last_run_info = {}
                if os.path.exists(LAST_RUN_FILE):
                    with open(LAST_RUN_FILE, "r", encoding="utf-8") as f:
                        last_run_info = json.load(f)
                self._send_json(200, {
                    "status": "ok",
                    "next_run_date": _next_run_4th(),
                    **last_run_info,
                })
            except Exception as e:
                self._send_json(500, {"status": "error", "message": str(e)})

        else:
            self.send_response(404)
            self.end_headers()


def run_server():
    os.makedirs(os.path.join(DIRECTORY, "logs"), exist_ok=True)
    os.chdir(DIRECTORY)
    _start_scheduler()

    with socketserver.TCPServer(("", PORT), Handler) as httpd:
        logger.info(f"서버 실행 중: http://localhost:{PORT}")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            pass
        finally:
            httpd.server_close()


if __name__ == "__main__":
    run_server()
